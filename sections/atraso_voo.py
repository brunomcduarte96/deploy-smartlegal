import streamlit as st
from utils.supabase_manager import SupabaseManager
from utils.audio_manager import AudioManager  # Vamos criar esse módulo
from datetime import datetime
import logging
import json
from openai import OpenAI
import os
from dotenv import load_dotenv
import time
from num2words import num2words
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
import io
from docx import Document
import re
from pathlib import Path
from google.oauth2 import service_account

logger = logging.getLogger(__name__)

def search_client(search_term, supabase):
    """Busca clientes pelo nome com autocomplete"""
    if not search_term:
        return []
    
    try:
        clients = supabase.search_clients_by_partial_name(search_term)
        return clients
    except Exception as e:
        logger.error(f"Erro ao buscar clientes: {str(e)}")
        return []

def render_client_section():
    """Renderiza a seção de seleção do cliente"""
    st.markdown("### 1. Selecionar Cliente")
    
    # Inicializar o SupabaseManager
    supabase = SupabaseManager()
    
    # Campo de busca do cliente
    search_term = st.text_input("Buscar cliente por nome:")
    
    if search_term:
        clients = search_client(search_term, supabase)
        if clients:
            # Criar lista de opções para o selectbox
            client_options = {
                f"{client['nome_completo']} - {client['cpf']}": client 
                for client in clients
            }
            
            # Dropdown para seleção do cliente
            selected_client = st.selectbox(
                "Selecione o cliente:",
                options=list(client_options.keys()),
                key="client_select"
            )
            
            if selected_client:
                # Salvar dados do cliente selecionado
                st.session_state.selected_client_data = client_options[selected_client]
                
                # Mostrar dados do cliente selecionado
                st.write("**Dados do cliente selecionado:**")
                
                # Criar duas colunas
                col1, col2 = st.columns(2)
                
                # Filtrar campos que queremos mostrar (excluindo 'id')
                campos = {k: v for k, v in st.session_state.selected_client_data.items() if k != 'id'}
                
                # Dividir os campos em duas listas de tamanho igual
                total_campos = len(campos)
                metade = (total_campos + 1) // 2  # +1 para garantir que todos os campos sejam mostrados
                
                # Preencher as colunas
                with col1:
                    for i, (key, value) in enumerate(campos.items()):
                        if i < metade:
                            st.write(f"**{key}:** {value}")
                
                with col2:
                    for i, (key, value) in enumerate(campos.items()):
                        if i >= metade:
                            st.write(f"**{key}:** {value}")
                
                # Buscar casos do cliente
                client_cases = supabase.get_client_cases(st.session_state.selected_client_data['id'])
                
                if client_cases:
                    st.markdown("---")
                    st.write("**Selecione o caso:**")
                    
                    # Criar lista de opções para o selectbox de casos
                    case_options = {
                        f"{case.get('chave_caso', 'Sem chave')} - {case.get('assunto_caso', 'Sem assunto')}": case 
                        for case in client_cases
                    }
                    
                    # Dropdown para seleção do caso
                    selected_case = st.selectbox(
                        "Casos do cliente",
                        options=list(case_options.keys()),
                        key="case_select"
                    )
                    
                    if selected_case:
                        # Salvar dados do caso selecionado
                        st.session_state.selected_case_data = case_options[selected_case]
                        
                        # Mostrar informações do caso
                        st.write("**Detalhes do caso:**")
                        case_col1, case_col2 = st.columns(2)
                        
                        with case_col1:
                            st.write(f"**Nome do Cliente:** {st.session_state.selected_case_data.get('nome_cliente', 'Não informado')}")
                            st.write(f"**Caso:** {st.session_state.selected_case_data.get('caso', 'Não informado')}")
                            st.write(f"**Assunto:** {st.session_state.selected_case_data.get('assunto_caso', 'Não informado')}")
                            st.write(f"**Responsável:** {st.session_state.selected_case_data.get('responsavel_comercial', 'Não informado')}")
                        
                        with case_col2:
                            st.write(f"**Chave do Caso:** {st.session_state.selected_case_data.get('chave_caso', 'Não informado')}")
                            st.write(f"**ID da Pasta:** {st.session_state.selected_case_data.get('pasta_caso_id', 'Não informado')}")
                            if pasta_url := st.session_state.selected_case_data.get('pasta_caso_url'):
                                st.write(f"**URL da Pasta:** [Acessar]({pasta_url})")
                            st.write(f"**Data de Criação:** {st.session_state.selected_case_data.get('created_at', 'Não informado')}")
                else:
                    st.info("Nenhum caso encontrado para este cliente")
        else:
            st.info("Nenhum cliente encontrado com esse nome.")

def render_company_section():
    """Renderiza a seção de seleção da empresa"""
    st.markdown("### 2. Selecionar Empresa")
    
    # Inicializar o SupabaseManager
    supabase = SupabaseManager()
    
    # Buscar todas as empresas
    companies = supabase.get_all_companies()
    
    if companies:
        # Criar lista de opções para o selectbox
        company_options = {
            f"{company['nome']} - {company['cnpj']}": company 
            for company in companies
        }
        
        # Dropdown para seleção da empresa
        selected_company = st.selectbox(
            "Selecione a empresa:",
            options=list(company_options.keys()),
            key="company_select"
        )
        
        if selected_company:
            # Salvar dados da empresa selecionada
            st.session_state.selected_company_data = company_options[selected_company]
            
            # Mostrar dados da empresa selecionada
            st.write("**Dados da empresa selecionada:**")
            for key, value in st.session_state.selected_company_data.items():
                st.write(f"**{key}:** {value}")
    else:
        st.error("Erro ao carregar lista de empresas.")

def get_openai_key():
    """Obtém a chave da API da OpenAI"""
    try:
        # Tentar ler a chave diretamente
        key = st.secrets.get("OPENAI_API_KEY")
        
        # Log para debug (não mostra a chave completa)
        if key:
            logger.info(f"Chave encontrada com tamanho: {len(key)}")
            logger.info(f"Primeiros caracteres: {key[:10]}")
        else:
            logger.error("Chave não encontrada em st.secrets")
            
        return key
        
    except Exception as e:
        logger.error(f"Erro ao obter chave da API: {str(e)}")
        raise Exception(f"Erro ao obter chave da API: {str(e)}")

def extract_flight_info(fatos_cliente):
    """Extrai informações do voo usando OpenAI"""
    try:
        api_key = get_openai_key()
        if not api_key:
            raise Exception("Chave da API não encontrada")
            
        client = OpenAI(
            api_key=api_key
        )
        
        # Teste simples antes de prosseguir
        test_response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[{"role": "user", "content": "test"}],
            max_tokens=5
        )
        logger.info("Teste de conexão bem sucedido")
        
        system_prompt = """Você é um assistente especializado em extrair informações sobre problemas com voos. 
        Analise o texto fornecido e extraia as informações solicitadas.
        Responda APENAS com um JSON contendo os campos solicitados.
        Se alguma informação não estiver disponível no texto, use "Não informado"."""
        
        user_prompt = f"""Extraia as seguintes informações do texto abaixo:
        - motivo_voo: Motivo da viagem
        - problema: Qual foi o problema enfrentado
        - compromisso_perdido: Se perdeu algum compromisso por conta do atraso
        - contexto: Condição física ou emocional do passageiro
        - momento_informacao: Quando foi informado do problema
        - tipo_voo: Nacional ou Internacional
        - origem_voo: Cidade/Aeroporto de origem
        - destino_voo: Cidade/Aeroporto de destino
        - escala: Se houve e onde foi
        - local_problema: Onde ocorreu o problema
        - data_voo_inicial: Data prevista do voo
        - horario_voo_inicial: Horário previsto do voo
        - data_voo_real: Data em que o voo de fato aconteceu
        - horario_voo_real: Horário em que o voo de fato aconteceu
        - tempo_atraso: Tempo de atraso (calcule se possível)
        - solicitou_reacomodacao: Se solicitou ser reacomodado em outro voo
        - opcao_reacomodacao: Descreva a opção de reacomodação recebida
        - recebeu_auxilio: Se recebeu algum auxílio como voucher de hotel ou comida
        - auxilio_recebido: Descreva o auxílio recebido
        - teve_custos: Se teve algum custo com uber, hotel ou comida
        - descricao_custos: Descreva os custos
        - valor_total_custos: Valor total dos custos (danos materiais)

        Texto: {fatos_cliente}"""

        try:
            response = client.chat.completions.create(
                model="gpt-4",
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt}
                ],
                temperature=0.3
            )
            
            # Converter a resposta para dict
            flight_info = json.loads(response.choices[0].message.content)
            return flight_info
            
        except Exception as e:
            logger.error(f"Erro na chamada da API OpenAI: {str(e)}")
            raise Exception("Erro ao processar o texto com a OpenAI. Por favor, tente novamente.")
            
    except Exception as e:
        logger.error(f"Erro ao processar texto com OpenAI: {str(e)}")
        raise e

def generate_facts():
    """Gera os fatos usando a API da OpenAI"""
    try:
        # Inicializar o cliente OpenAI
        client = OpenAI(
            api_key=st.secrets["OPENAI_API_KEY"]
        )
        
        # Verificar se flight_info existe no session_state
        if not hasattr(st.session_state, 'flight_info'):
            raise Exception("Informações do voo não encontradas. Por favor, preencha os dados do voo primeiro.")

        # Obter o texto dos fatos relatados
        fatos_relatados = st.session_state.get('transcription', '')

        # Criar o dicionário flight_info com todos os campos
        flight_info = {
            'tipo_voo': st.session_state.flight_info.get('tipo_voo', 'Não informado'),
            'origem_voo': st.session_state.flight_info.get('origem_voo', 'Não informado'),
            'destino_voo': st.session_state.flight_info.get('destino_voo', 'Não informado'),
            'escala': st.session_state.flight_info.get('escala', 'Não informado'),
            'data_voo_inicial': st.session_state.flight_info.get('data_voo_inicial', 'Não informado'),
            'horario_voo_inicial': st.session_state.flight_info.get('horario_voo_inicial', 'Não informado'),
            'data_voo_real': st.session_state.flight_info.get('data_voo_real', 'Não informado'),
            'horario_voo_real': st.session_state.flight_info.get('horario_voo_real', 'Não informado'),
            'tempo_atraso': st.session_state.flight_info.get('tempo_atraso', 'Não informado'),
            'motivo_voo': st.session_state.flight_info.get('motivo_voo', 'Não informado'),
            'problema': st.session_state.flight_info.get('problema', 'Não informado'),
            'local_problema': st.session_state.flight_info.get('local_problema', 'Não informado'),
            'momento_informacao': st.session_state.flight_info.get('momento_informacao', 'Não informado'),
            'compromisso_perdido': st.session_state.flight_info.get('compromisso_perdido', 'Não informado'),
            'contexto': st.session_state.flight_info.get('contexto', 'Não informado'),
            'solicitou_reacomodacao': st.session_state.flight_info.get('solicitou_reacomodacao', 'Não informado'),
            'opcao_reacomodacao': st.session_state.flight_info.get('opcao_reacomodacao', 'Não informado'),
            'recebeu_auxilio': st.session_state.flight_info.get('recebeu_auxilio', 'Não informado'),
            'auxilio_recebido': st.session_state.flight_info.get('auxilio_recebido', 'Não informado'),
            'teve_custos': st.session_state.flight_info.get('teve_custos', 'Não informado'),
            'descricao_custos': st.session_state.flight_info.get('descricao_custos', 'Não informado'),
            'valor_total_custos': st.session_state.flight_info.get('valor_total_custos', 'Não informado')
        }
        
        message_content = f"""
        FATOS RELATADOS PELO CLIENTE:
        {fatos_relatados}
        
        INFORMAÇÕES DO VOO:
        
        DADOS DO VOO:
        - Tipo: {flight_info['tipo_voo']}
        - Origem: {flight_info['origem_voo']}
        - Destino: {flight_info['destino_voo']}
        - Escala: {flight_info['escala']}
        - Data Prevista: {flight_info['data_voo_inicial']}
        - Horário Previsto: {flight_info['horario_voo_inicial']}
        - Data Real: {flight_info['data_voo_real']}
        - Horário Real: {flight_info['horario_voo_real']}
        - Tempo de Atraso: {flight_info['tempo_atraso']}
        - Motivo da Viagem: {flight_info['motivo_voo']}
        
        DETALHES DO PROBLEMA:
        - Problema: {flight_info['problema']}
        - Local do Problema: {flight_info['local_problema']}
        - Momento da Informação: {flight_info['momento_informacao']}
        - Compromisso Perdido: {flight_info['compromisso_perdido']}
        - Contexto do Passageiro: {flight_info['contexto']}
        
        AUXÍLIOS E CUSTOS:
        - Solicitou Reacomodação: {flight_info['solicitou_reacomodacao']}
        - Opção Recebida: {flight_info['opcao_reacomodacao']}
        - Recebeu Auxílio: {flight_info['recebeu_auxilio']}
        - Auxílio Recebido: {flight_info['auxilio_recebido']}
        - Teve Custos: {flight_info['teve_custos']}
        - Descrição dos Custos: {flight_info['descricao_custos']}
        - Valor Total: {flight_info['valor_total_custos']}
        """

        # Criar um thread
        thread = client.beta.threads.create()

        # Adicionar a mensagem ao thread
        message = client.beta.threads.messages.create(
            thread_id=thread.id,
            role="user",
            content=message_content
        )

        # Executar o assistente
        run = client.beta.threads.runs.create(
            thread_id=thread.id,
            assistant_id="asst_Yjlfjp23PwM2qiN6ckjxUspB"
        )

        # Aguardar a conclusão
        while True:
            run_status = client.beta.threads.runs.retrieve(
                thread_id=thread.id,
                run_id=run.id
            )
            if run_status.status == 'completed':
                break
            time.sleep(1)  # Esperar 1 segundo antes de verificar novamente

        # Obter a resposta
        messages = client.beta.threads.messages.list(
            thread_id=thread.id
        )
        
        # A resposta mais recente é a primeira da lista
        generated_facts = messages.data[0].content[0].text.value

        return generated_facts

    except Exception as e:
        logger.error(f"Erro ao gerar fatos: {str(e)}")
        raise Exception(f"Erro ao gerar fatos: {str(e)}")

def render_facts_section():
    """Renderiza a seção de fatos do voo"""
    st.markdown("---")
    st.subheader("3. Fatos do Voo")
    
    audio_manager = AudioManager()
    
    # Botão para abrir o conversor de áudio online
    st.markdown("""
        <a href="https://online-audio-converter.com/pt/" target="_blank">
            <button style="
                background-color: #4CAF50;
                border: none;
                color: white;
                padding: 10px 24px;
                text-align: center;
                text-decoration: none;
                display: inline-block;
                font-size: 16px;
                margin: 4px 2px;
                cursor: pointer;
                border-radius: 4px;">
                Converter Áudio
            </button>
        </a>
    """, unsafe_allow_html=True)
    
    st.write("") # Espaço entre os elementos
    
    # Upload do arquivo de áudio
    audio_file = st.file_uploader(
        "Fazer upload do áudio",
        type=['wav'],
        help="Por favor, envie o áudio no formato WAV"
    )
    
    if audio_file:
        # Salvar arquivo temporariamente
        audio_path = audio_manager.save_temp_audio(audio_file)
        
        # Botão para transcrever
        if st.button("Transcrever Áudio"):
            with st.spinner("Transcrevendo áudio..."):
                try:
                    transcription = audio_manager.transcribe_audio(audio_path)
                    st.session_state.transcription = transcription
                    st.session_state.last_text = transcription  # Guardar último texto processado
                except Exception as e:
                    logger.error(f"Erro na transcrição: {str(e)}")
                    st.error("Erro ao transcrever o áudio")
    
    # Campo editável com a transcrição
    fatos_cliente = st.text_area(
        "Fatos relatados pelo cliente",
        value=st.session_state.get('transcription', ''),
        height=300,
        key="fatos_cliente"
    )
    
    # Atualizar variável na session_state quando o texto for editado
    if fatos_cliente != st.session_state.get('transcription', ''):
        st.session_state.transcription = fatos_cliente
    
    # Botão para reconhecer informações
    if st.button("Reconhecer Informações", type="primary"):
        if not fatos_cliente:
            st.error("Por favor, forneça os fatos do voo antes de reconhecer as informações.")
            return
            
        try:
            with st.spinner("Analisando informações..."):
                flight_info = extract_flight_info(fatos_cliente)
                st.session_state.flight_info = flight_info
                st.success("Informações reconhecidas com sucesso!")
                st.rerun()  # Atualizar a página para mostrar as informações
        except Exception as e:
            st.error(f"Erro ao processar as informações: {str(e)}")
    
    # Inicializar flight_info se não existir
    if 'flight_info' not in st.session_state:
        st.session_state.flight_info = {
            'tipo_voo': '', 'origem_voo': '', 'destino_voo': '', 'escala': '',
            'data_voo_inicial': '', 'horario_voo_inicial': '', 'data_voo_real': '',
            'horario_voo_real': '', 'tempo_atraso': '', 'motivo_voo': '',
            'problema': '', 'local_problema': '', 'momento_informacao': '',
            'compromisso_perdido': '', 'contexto': '', 'solicitou_reacomodacao': '',
            'opcao_reacomodacao': '', 'recebeu_auxilio': '', 'auxilio_recebido': '',
            'teve_custos': '', 'descricao_custos': '', 'valor_total_custos': ''
        }
    
    # Mostrar campos editáveis
    st.subheader("Informações do Voo:")
    
    col1, col2, col3 = st.columns(3)
    
    with col1:
        st.write("**Dados do Voo:**")
        st.session_state.flight_info['tipo_voo'] = st.text_input(
            "Tipo de Voo", 
            value=st.session_state.flight_info.get('tipo_voo', '')
        )
        st.session_state.flight_info['origem_voo'] = st.text_input(
            "Origem", 
            value=st.session_state.flight_info.get('origem_voo', '')
        )
        st.session_state.flight_info['destino_voo'] = st.text_input(
            "Destino", 
            value=st.session_state.flight_info.get('destino_voo', '')
        )
        st.session_state.flight_info['escala'] = st.text_input(
            "Escala", 
            value=st.session_state.flight_info.get('escala', '')
        )
        
        # Campo de data prevista
        try:
            data_prevista = st.date_input(
                "Data Prevista",
                value=None if not st.session_state.flight_info.get('data_voo_inicial') or 
                          st.session_state.flight_info['data_voo_inicial'] == "Não informado" or
                          format_date(st.session_state.flight_info['data_voo_inicial']) == ""
                else datetime.strptime(format_date(st.session_state.flight_info['data_voo_inicial']), "%d/%m/%Y"),
                format="DD/MM/YYYY"
            )
            st.session_state.flight_info['data_voo_inicial'] = data_prevista.strftime("%d/%m/%Y") if data_prevista else ""
        except:
            data_prevista = st.date_input("Data Prevista", value=None, format="DD/MM/YYYY")
            st.session_state.flight_info['data_voo_inicial'] = data_prevista.strftime("%d/%m/%Y") if data_prevista else ""
        
        # Campo de horário previsto
        try:
            horario_previsto = st.time_input(
                "Horário Previsto",
                value=None if not st.session_state.flight_info.get('horario_voo_inicial') or
                          st.session_state.flight_info['horario_voo_inicial'] == "Não informado" or
                          format_time(st.session_state.flight_info['horario_voo_inicial']) == ""
                else datetime.strptime(format_time(st.session_state.flight_info['horario_voo_inicial']), "%H:%M").time()
            )
            st.session_state.flight_info['horario_voo_inicial'] = horario_previsto.strftime("%H:%M") if horario_previsto else ""
        except:
            horario_previsto = st.time_input("Horário Previsto", value=None)
            st.session_state.flight_info['horario_voo_inicial'] = horario_previsto.strftime("%H:%M") if horario_previsto else ""
        
        # Campo de data real
        try:
            data_real = st.date_input(
                "Data Real",
                value=None if not st.session_state.flight_info.get('data_voo_real') or
                          st.session_state.flight_info['data_voo_real'] == "Não informado" or
                          format_date(st.session_state.flight_info['data_voo_real']) == ""
                else datetime.strptime(format_date(st.session_state.flight_info['data_voo_real']), "%d/%m/%Y"),
                format="DD/MM/YYYY"
            )
            st.session_state.flight_info['data_voo_real'] = data_real.strftime("%d/%m/%Y") if data_real else ""
        except:
            data_real = st.date_input("Data Real", value=None, format="DD/MM/YYYY")
            st.session_state.flight_info['data_voo_real'] = data_real.strftime("%d/%m/%Y") if data_real else ""
        
        # Campo de horário real
        try:
            horario_real = st.time_input(
                "Horário Real",
                value=None if not st.session_state.flight_info.get('horario_voo_real') or
                          st.session_state.flight_info['horario_voo_real'] == "Não informado" or
                          format_time(st.session_state.flight_info['horario_voo_real']) == ""
                else datetime.strptime(format_time(st.session_state.flight_info['horario_voo_real']), "%H:%M").time()
            )
            st.session_state.flight_info['horario_voo_real'] = horario_real.strftime("%H:%M") if horario_real else ""
        except:
            horario_real = st.time_input("Horário Real", value=None)
            st.session_state.flight_info['horario_voo_real'] = horario_real.strftime("%H:%M") if horario_real else ""
        
        st.session_state.flight_info['tempo_atraso'] = st.text_input(
            "Tempo de Atraso", 
            value=st.session_state.flight_info.get('tempo_atraso', '')
        )
    
    with col2:
        st.write("**Detalhes do Problema:**")
        st.session_state.flight_info['motivo_voo'] = st.text_input(
            "Motivo da Viagem", 
            value=st.session_state.flight_info.get('motivo_voo', '')
        )
        st.session_state.flight_info['problema'] = st.text_area(
            "Problema Enfrentado", 
            value=st.session_state.flight_info.get('problema', ''),
            height=100
        )
        st.session_state.flight_info['local_problema'] = st.text_input(
            "Local do Problema", 
            value=st.session_state.flight_info.get('local_problema', '')
        )
        st.session_state.flight_info['momento_informacao'] = st.text_input(
            "Momento da Informação", 
            value=st.session_state.flight_info.get('momento_informacao', '')
        )
        st.session_state.flight_info['compromisso_perdido'] = st.text_area(
            "Compromisso Perdido", 
            value=st.session_state.flight_info.get('compromisso_perdido', ''),
            height=100
        )
        st.session_state.flight_info['contexto'] = st.text_area(
            "Contexto do Passageiro", 
            value=st.session_state.flight_info.get('contexto', ''),
            height=100
        )
    
    with col3:
        st.write("**Auxílios e Custos:**")
        st.session_state.flight_info['solicitou_reacomodacao'] = st.text_input(
            "Solicitou Reacomodação?", 
            value=st.session_state.flight_info.get('solicitou_reacomodacao', '')
        )
        st.session_state.flight_info['opcao_reacomodacao'] = st.text_area(
            "Opção de Reacomodação Recebida", 
            value=st.session_state.flight_info.get('opcao_reacomodacao', ''),
            height=100
        )
        st.session_state.flight_info['recebeu_auxilio'] = st.text_input(
            "Recebeu Auxílio?", 
            value=st.session_state.flight_info.get('recebeu_auxilio', '')
        )
        st.session_state.flight_info['auxilio_recebido'] = st.text_area(
            "Auxílio Recebido", 
            value=st.session_state.flight_info.get('auxilio_recebido', ''),
            height=100
        )
        st.session_state.flight_info['teve_custos'] = st.text_input(
            "Teve Custos Extras?", 
            value=st.session_state.flight_info.get('teve_custos', '')
        )
        st.session_state.flight_info['descricao_custos'] = st.text_area(
            "Descrição dos Custos", 
            value=st.session_state.flight_info.get('descricao_custos', ''),
            height=100
        )
        st.session_state.flight_info['valor_total_custos'] = st.text_input(
            "Valor Total dos Custos", 
            value=st.session_state.flight_info.get('valor_total_custos', '')
        )
    
    # Seção de Geração de Fatos
    st.markdown("---")
    
    # Verificar se todos os dados necessários estão presentes
    if st.button("Gerar Fatos", type="primary"):
        if not fatos_cliente:
            st.error("Por favor, preencha os fatos relatados pelo cliente antes de gerar os fatos.")
            return
            
        if not hasattr(st.session_state, 'flight_info'):
            st.error("Por favor, preencha as informações do voo antes de gerar os fatos.")
            return
            
        try:
            with st.spinner("Gerando fatos do caso..."):
                generated_facts = generate_facts()
                st.session_state.generated_facts = generated_facts
                st.success("Fatos gerados com sucesso!")
        except Exception as e:
            st.error(f"Erro ao gerar fatos: {str(e)}")
            if st.button("Tentar Novamente"):
                st.rerun()
    
    # Mostrar os fatos gerados em uma caixa editável se existirem
    if 'generated_facts' in st.session_state:
        st.markdown("### Fatos Gerados")
        st.markdown("_Você pode editar o texto abaixo se necessário:_")
        edited_facts = st.text_area(
            "Fatos do Caso",
            value=st.session_state.generated_facts,
            height=400,
            key="facts_text_area"
        )
        
        # Salvar os fatos editados na session_state
        if edited_facts != st.session_state.generated_facts:
            st.session_state.generated_facts = edited_facts
            st.info("As alterações foram salvas!")

    st.markdown("---")
    
    # Botões de ação (sempre visíveis)
    st.markdown("### Ações")
    
    col1, col2, col3 = st.columns([1, 2, 1])
    
    with col2:
        # Botão Gerar Novamente
        if st.button("Gerar Novamente", type="primary", use_container_width=True, key="btn_gerar"):
            if not fatos_cliente:
                st.error("Por favor, preencha os fatos relatados pelo cliente antes de gerar os fatos.")
                return
            
            if not hasattr(st.session_state, 'flight_info'):
                st.error("Por favor, preencha as informações do voo antes de gerar os fatos.")
                return
            
            try:
                with st.spinner("Gerando novos fatos..."):
                    generated_facts = generate_facts()
                    st.session_state.generated_facts = generated_facts
                    st.rerun()
            except Exception as e:
                st.error(f"Erro ao gerar novos fatos: {str(e)}")
        
        # Botão Salvar para Treinamento
        if st.button("Salvar para Treinamento", type="primary", use_container_width=True, key="btn_salvar"):
            if not hasattr(st.session_state, 'generated_facts'):
                st.error("Por favor, gere os fatos antes de salvar para treinamento.")
                return
                
            try:
                supabase = SupabaseManager()
                caso = st.session_state.selected_case_data.get('assunto_caso', 'Não informado')
                
                # Recriar o message_content
                message_content = f"""
                FATOS RELATADOS PELO CLIENTE:
                {fatos_cliente}
                
                DADOS COMPLEMENTARES DO VOO:
                
                DADOS DO VOO:
                - Tipo: {st.session_state.flight_info['tipo_voo']}
                - Origem: {st.session_state.flight_info['origem_voo']}
                - Destino: {st.session_state.flight_info['destino_voo']}
                - Escala: {st.session_state.flight_info['escala']}
                - Data Prevista: {st.session_state.flight_info['data_voo_inicial']}
                - Horário Previsto: {st.session_state.flight_info['horario_voo_inicial']}
                - Data Real: {st.session_state.flight_info['data_voo_real']}
                - Horário Real: {st.session_state.flight_info['horario_voo_real']}
                - Tempo de Atraso: {st.session_state.flight_info['tempo_atraso']}
                
                DETALHES DO PROBLEMA:
                - Motivo da Viagem: {st.session_state.flight_info['motivo_voo']}
                - Problema: {st.session_state.flight_info['problema']}
                - Local do Problema: {st.session_state.flight_info['local_problema']}
                - Momento da Informação: {st.session_state.flight_info['momento_informacao']}
                - Compromisso Perdido: {st.session_state.flight_info['compromisso_perdido']}
                - Contexto do Passageiro: {st.session_state.flight_info['contexto']}
                
                AUXÍLIOS E CUSTOS:
                - Solicitou Reacomodação: {st.session_state.flight_info['solicitou_reacomodacao']}
                - Opção Recebida: {st.session_state.flight_info['opcao_reacomodacao']}
                - Recebeu Auxílio: {st.session_state.flight_info['recebeu_auxilio']}
                - Auxílio Recebido: {st.session_state.flight_info['auxilio_recebido']}
                - Teve Custos: {st.session_state.flight_info['teve_custos']}
                - Descrição dos Custos: {st.session_state.flight_info['descricao_custos']}
                - Valor Total: {st.session_state.flight_info['valor_total_custos']}
                """
                
                # Salvar no banco de dados
                supabase.save_facts_for_training(
                    caso=caso,
                    input_text=message_content,
                    output_text=st.session_state.generated_facts
                )
                
                st.success("Fatos salvos para treinamento com sucesso!")
                
            except Exception as e:
                st.error(f"Erro ao salvar para treinamento: {str(e)}")
    
    st.markdown("---")

    # Campo para Vara Cível
    st.markdown("### 4. Vara Cível")
    vara_civil = st.text_input("Vara Cível", key="vara_civil")

    # Seção Jurisprudência Seção Deveres do Transportador
    st.markdown("### 5. Jurisprudência Seção Deveres do Transportador")
    
    try:
        supabase = SupabaseManager()
        jurisprudencias = supabase.get_jurisprudencias_aereo()
        
        if jurisprudencias:
            # Criar dicionário com nome como chave e dados completos como valor
            jurisprudencia_options = {
                jurisprudencia['nome']: jurisprudencia 
                for jurisprudencia in jurisprudencias
            }
            
            # Definir o valor padrão
            default_option = "cancelamento_voo_sem_aviso_indenizacao_10k"
            default_index = list(jurisprudencia_options.keys()).index(default_option) if default_option in jurisprudencia_options else 0
            
            # Dropdown para seleção da jurisprudência
            selected_jurisprudencia = st.selectbox(
                "Selecione a Jurisprudência - Seção Deveres do Transportador",
                options=list(jurisprudencia_options.keys()),
                index=default_index,
                key="jurisprudencia_select"
            )
            
            if selected_jurisprudencia:
                # Mostrar detalhes da jurisprudência selecionada
                jurisprudencia_data = jurisprudencia_options[selected_jurisprudencia]
                
                # Salvar no session_state
                st.session_state['tribunal_jurisprudencia_deveres_transportador'] = jurisprudencia_data['Tribunal']
                st.session_state['jurisprudencia_deveres_transportador'] = jurisprudencia_data['texto']
                
                # Mostrar tribunal e texto um embaixo do outro
                st.write(f"**Tribunal:** {jurisprudencia_data['Tribunal']}")
                st.write("**Texto:**")
                st.markdown(jurisprudencia_data['texto'])
        else:
            st.info("Nenhuma jurisprudência encontrada")
            
    except Exception as e:
        logger.error(f"Erro ao carregar jurisprudências: {str(e)}")
        st.error("Erro ao carregar lista de jurisprudências")

    # Seção Jurisprudência Seção Da Inteligência
    st.markdown("### 6. Jurisprudência Seção Da Inteligência")
    
    try:
        if jurisprudencias:
            # Definir o valor padrão
            default_option = "atraso_conexao_internacional_24h_indenizacao_10k"
            default_index = list(jurisprudencia_options.keys()).index(default_option) if default_option in jurisprudencia_options else 0
            
            # Dropdown para seleção da jurisprudência
            selected_jurisprudencia_intel = st.selectbox(
                "Selecione a Jurisprudência - Seção Da Inteligência",
                options=list(jurisprudencia_options.keys()),
                index=default_index,
                key="jurisprudencia_intel_select"
            )
            
            if selected_jurisprudencia_intel:
                # Mostrar detalhes da jurisprudência selecionada
                jurisprudencia_data_intel = jurisprudencia_options[selected_jurisprudencia_intel]
                
                # Salvar no session_state
                st.session_state['tribunal_jurisprudencia_da_inteligencia'] = jurisprudencia_data_intel['Tribunal']
                st.session_state['jurisprudencia_da_inteligencia'] = jurisprudencia_data_intel['texto']
                
                # Mostrar tribunal e texto um embaixo do outro
                st.write(f"**Tribunal:** {jurisprudencia_data_intel['Tribunal']}")
                st.write("**Texto:**")
                st.markdown(jurisprudencia_data_intel['texto'])
        else:
            st.info("Nenhuma jurisprudência encontrada")
            
    except Exception as e:
        logger.error(f"Erro ao carregar jurisprudências: {str(e)}")
        st.error("Erro ao carregar lista de jurisprudências")

    # Seção Jurisprudência Seção Da Responsabilidade
    st.markdown("### 7. Jurisprudência Seção Da Responsabilidade")
    
    try:
        if jurisprudencias:
            # Definir o valor padrão
            default_option = "atraso_voo_ma_assistencia_30h_indenizacao_15k"
            default_index = list(jurisprudencia_options.keys()).index(default_option) if default_option in jurisprudencia_options else 0
            
            # Dropdown para seleção da jurisprudência
            selected_jurisprudencia_resp = st.selectbox(
                "Selecione a Jurisprudência - Seção Da Responsabilidade",
                options=list(jurisprudencia_options.keys()),
                index=default_index,
                key="jurisprudencia_resp_select"
            )
            
            if selected_jurisprudencia_resp:
                # Mostrar detalhes da jurisprudência selecionada
                jurisprudencia_data_resp = jurisprudencia_options[selected_jurisprudencia_resp]
                
                # Salvar no session_state
                st.session_state['tribunal_jurisprudencia_da_responsabilidadea'] = jurisprudencia_data_resp['Tribunal']
                st.session_state['jurisprudencia_da_responsabilidadea'] = jurisprudencia_data_resp['texto']
                
                # Mostrar tribunal e texto um embaixo do outro
                st.write(f"**Tribunal:** {jurisprudencia_data_resp['Tribunal']}")
                st.write("**Texto:**")
                st.markdown(jurisprudencia_data_resp['texto'])
        else:
            st.info("Nenhuma jurisprudência encontrada")
            
    except Exception as e:
        logger.error(f"Erro ao carregar jurisprudências: {str(e)}")
        st.error("Erro ao carregar lista de jurisprudências")

    st.markdown("---")
    
    # Seção Jurisprudência Seção Dos Prejuízos
    st.markdown("### 8. Jurisprudência Seção Dos Prejuízos")
    
    try:
        if jurisprudencias:
            # Definir o valor padrão
            default_option = "cancelamento_voo_internacional_24h_indenizacao_10k"
            default_index = list(jurisprudencia_options.keys()).index(default_option) if default_option in jurisprudencia_options else 0
            
            # Dropdown para seleção da jurisprudência
            selected_jurisprudencia_prej = st.selectbox(
                "Selecione a Jurisprudência - Seção Dos Prejuízos",
                options=list(jurisprudencia_options.keys()),
                index=default_index,
                key="jurisprudencia_prej_select"
            )
            
            if selected_jurisprudencia_prej:
                # Mostrar detalhes da jurisprudência selecionada
                jurisprudencia_data_prej = jurisprudencia_options[selected_jurisprudencia_prej]
                
                # Salvar no session_state
                st.session_state['tribunal_jurisprudencia_dos_prejuizos'] = jurisprudencia_data_prej['Tribunal']
                st.session_state['jurisprudencia_dos_prejuizos'] = jurisprudencia_data_prej['texto']
                
                # Mostrar tribunal e texto um embaixo do outro
                st.write(f"**Tribunal:** {jurisprudencia_data_prej['Tribunal']}")
                st.write("**Texto:**")
                st.markdown(jurisprudencia_data_prej['texto'])
        else:
            st.info("Nenhuma jurisprudência encontrada")
            
    except Exception as e:
        logger.error(f"Erro ao carregar jurisprudências: {str(e)}")
        st.error("Erro ao carregar lista de jurisprudências")

    st.markdown("---")
    
    # Seção Dano Moral e Dano Material
    st.markdown("### 9. Dano Moral e Dano Material")
    
    # 9.1 - Valor Danos Morais
    st.markdown("#### 9.1. Valor Danos Morais")
    
    def on_valor_danos_morais_change():
        """Callback para atualizar os valores formatados quando o input mudar"""
        valor = st.session_state.valor_danos_morais
        # Salvar valor formatado
        st.session_state['valor_dano_moral'] = f"R$ {valor:.2f}"
        # Converter para extenso
        try:
            valor_extenso = num2words(valor, lang='pt_BR', to='currency')
            valor_extenso = valor_extenso.capitalize()
            if valor != 1:
                valor_extenso = valor_extenso.replace(" real", " reais")
            st.session_state['valor_dano_moral_extenso'] = valor_extenso
        except Exception as e:
            logger.error(f"Erro ao converter valor para extenso: {str(e)}")
    
    valor_danos_morais = st.number_input(
        "Valor dos Danos Morais (R$)",
        min_value=0.0,
        format="%.2f",
        step=100.0,
        key="valor_danos_morais",
        on_change=on_valor_danos_morais_change  # Adicionar callback
    )
    
    # Mostrar valor por extenso se disponível
    if 'valor_dano_moral_extenso' in st.session_state:
        st.write(f"**Valor por extenso:** {st.session_state.valor_dano_moral_extenso}")
    
    # 9.2 - Motivos Danos Morais
    st.markdown("#### 9.2. Motivos Danos Morais")
    motivos_exemplo = """- O cancelamento do voo do Autor pela empresa Ré após longa jornada no aeroporto;

- A recusa da empresa Ré em reacomodar o Autor em um voo que tivesse o horário de partida próximo ao anterior, realocando o autor para um voo só no dia seguinte, fazendo ele chegar ao destino final 15 horas depois do previsto; 

- O atraso substancial de 15 horas na sua viagem;

- A falta de prestação de assistência material pela Empresa Ré ao Autor."""

    motivos_danos_morais = st.text_area(
        "Motivos dos Danos Morais",
        value=motivos_exemplo,
        height=300,
        key="motivos_danos_morais"
    )
    
    # 9.3 - Danos Materiais
    st.markdown("#### 9.3. Danos Materiais (já preenchido na seção 3)")
    
    # Obter o valor dos custos da seção 3
    valor_danos_materiais = 0.0
    if hasattr(st.session_state, 'flight_info'):
        valor_str = st.session_state.flight_info.get('valor_total_custos', '0')
        # Remover 'R$' e outros caracteres não numéricos, exceto ponto e vírgula
        valor_str = ''.join(c for c in valor_str if c.isdigit() or c in '.,')
        try:
            # Substituir vírgula por ponto e converter para float
            valor_danos_materiais = float(valor_str.replace(',', '.'))
        except:
            valor_danos_materiais = 0.0
    
    st.write(f"**Valor dos Danos Materiais:** R$ {valor_danos_materiais:.2f}")
    
    # Converter o valor para extenso
    try:
        valor_materiais_extenso = num2words(valor_danos_materiais, lang='pt_BR', to='currency')
        valor_materiais_extenso = valor_materiais_extenso.capitalize()
        if valor_danos_materiais != 1:
            valor_materiais_extenso = valor_materiais_extenso.replace(" real", " reais")
        st.session_state['valor_danos_material_extenso'] = valor_materiais_extenso  # Salvar no session_state
        st.write(f"**Valor por extenso:** {valor_materiais_extenso}")
    except Exception as e:
        logger.error(f"Erro ao converter valor material para extenso: {str(e)}")
        st.error("Erro ao converter valor material para extenso")
    
    # 9.4 - Valor Total da Causa
    st.markdown("#### 9.4. Valor Total da Causa")
    
    # Calcular valor total (danos materiais + morais)
    valor_total_causa = valor_danos_materiais + valor_danos_morais
    
    st.write(f"**Valor Total da Causa:** R$ {valor_total_causa:.2f}")
    
    # Converter o valor total para extenso
    try:
        valor_total_extenso = num2words(valor_total_causa, lang='pt_BR', to='currency')
        valor_total_extenso = valor_total_extenso.capitalize()
        if valor_total_causa != 1:
            valor_total_extenso = valor_total_extenso.replace(" real", " reais")
        st.session_state['valor_dano_moral_material_extenso'] = valor_total_extenso  # Salvar no session_state
        st.session_state['valor_dano_moral_material'] = f"R$ {valor_total_causa:.2f}"  # Salvar valor formatado
        st.write(f"**Valor por extenso:** {valor_total_extenso}")
    except Exception as e:
        logger.error(f"Erro ao converter valor total para extenso: {str(e)}")
        st.error("Erro ao converter valor total para extenso")
    
    st.markdown("---")
    
    # Seção 10. Data
    st.markdown("### 10. Data")
    
    # Obter data atual
    data_atual = datetime.now()
    
    # Formato DD-MM-YYYY
    data_formatada = data_atual.strftime("%d-%m-%Y")
    st.write(f"**Data:** {data_formatada}")
    
    # Formato por extenso (25 de março de 2025)
    meses = {
        1: 'janeiro', 2: 'fevereiro', 3: 'março', 4: 'abril',
        5: 'maio', 6: 'junho', 7: 'julho', 8: 'agosto',
        9: 'setembro', 10: 'outubro', 11: 'novembro', 12: 'dezembro'
    }
    
    data_extenso = f"{data_atual.day} de {meses[data_atual.month]} de {data_atual.year}"
    st.write(f"**Data por extenso:** {data_extenso}")
    st.session_state['data_extenso'] = data_extenso  # Salvar no session_state
    
    st.markdown("---")
    
    # Botão para gerar a petição completa
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        if st.button("Gerar Petição Completa", type="primary", use_container_width=True):
            try:
                with st.spinner("Gerando petição inicial..."):
                    doc_link = generate_and_save_petition(st.session_state)
                    st.success("Petição gerada com sucesso!")
                    st.markdown(f"[Clique aqui para visualizar a petição]({doc_link})")
            except Exception as e:
                st.error(f"Erro ao gerar petição: {str(e)}")
    
    # Seção de verificação dos campos
    if st.checkbox("Mostrar dados que serão usados na petição"):
        st.markdown("### Verificação dos Dados da Petição")
        
        # Criar um dicionário com todos os campos
        campos = {
            'nome_completo': st.session_state.selected_client_data.get('nome_completo', ''),
            'nacionalidade': st.session_state.selected_client_data.get('nacionalidade', ''),
            'estado_civil': st.session_state.selected_client_data.get('estado_civil', ''),
            'profissao': st.session_state.selected_client_data.get('profissao', ''),
            'rg': st.session_state.selected_client_data.get('rg', ''),
            'cpf': st.session_state.selected_client_data.get('cpf', ''),
            'endereco': st.session_state.selected_client_data.get('endereco', ''),
            'bairro': st.session_state.selected_client_data.get('bairro', ''),
            'cidade': st.session_state.selected_client_data.get('cidade', ''),
            'cep': st.session_state.selected_client_data.get('cep', ''),
            'nome_empresa_re': st.session_state.selected_company_data.get('nome', ''),
            'cnpj_empresa_re': st.session_state.selected_company_data.get('cnpj', ''),
            'endereco_empresa_re': st.session_state.selected_company_data.get('endereco', ''),
            'dos_fatos': st.session_state.get('generated_facts', ''),
            'tempo_atraso': st.session_state.flight_info.get('tempo_atraso', ''),
            'valor_danos_material': st.session_state.flight_info.get('valor_total_custos', ''),
            'valor_danos_material_extenso': st.session_state.get('valor_danos_material_extenso', ''),
            'explicacao_danos_material': st.session_state.flight_info.get('descricao_custos', ''),
            'vara_civil': st.session_state.get('vara_civil', ''),
            'tribunal_jurisprudencia_deveres_transportador': st.session_state.get('tribunal_jurisprudencia_deveres_transportador', ''),
            'jurisprudencia_deveres_transportador': st.session_state.get('jurisprudencia_deveres_transportador', ''),
            'tribunal_jurisprudencia_da_inteligencia': st.session_state.get('tribunal_jurisprudencia_da_inteligencia', ''),
            'jurisprudencia_da_inteligencia': st.session_state.get('jurisprudencia_da_inteligencia', ''),
            'tribunal_jurisprudencia_da_responsabilidadea': st.session_state.get('tribunal_jurisprudencia_da_responsabilidadea', ''),
            'jurisprudencia_da_responsabilidadea': st.session_state.get('jurisprudencia_da_responsabilidadea', ''),
            'tribunal_jurisprudencia_dos_prejuizos': st.session_state.get('tribunal_jurisprudencia_dos_prejuizos', ''),
            'jurisprudencia_dos_prejuizos': st.session_state.get('jurisprudencia_dos_prejuizos', ''),
            'motivos_danos_moral': st.session_state.get('motivos_danos_morais', ''),
            'valor_dano_moral': st.session_state.get('valor_dano_moral', ''),
            'valor_dano_moral_extenso': st.session_state.get('valor_dano_moral_extenso', ''),
            'valor_dano_moral_material': st.session_state.get('valor_dano_moral_material', ''),
            'valor_dano_moral_material_extenso': st.session_state.get('valor_dano_moral_material_extenso', ''),
            'data_extenso': st.session_state.get('data_extenso', '')
        }
        
        # Mostrar cada campo em um expander
        for key, value in campos.items():
            with st.expander(f"{{{{ {key} }}}}"):
                st.write(value)
    
    st.markdown("---")

def format_date(date_str):
    """Converte a data para o formato dd/mm/aaaa"""
    if not date_str or date_str == "Não informado":
        return ""
    try:
        # Remover possíveis textos extras
        date_str = date_str.lower().replace(" de ", "/")
        
        # Converter nomes dos meses para números
        meses = {
            'janeiro': '01', 'fevereiro': '02', 'março': '03', 'abril': '04',
            'maio': '05', 'junho': '06', 'julho': '07', 'agosto': '08',
            'setembro': '09', 'outubro': '10', 'novembro': '11', 'dezembro': '12',
            'marco': '03'  # Versão sem acento
        }
        
        for mes, num in meses.items():
            date_str = date_str.replace(mes, num)
        
        # Tentar diferentes formatos de entrada
        formatos = [
            "%d/%m/%Y",  # 31/12/2023
            "%Y-%m-%d",  # 2023-12-31
            "%d-%m-%Y",  # 31-12-2023
            "%d/%m/%y",  # 31/12/23
            "%d %m %Y",  # 31 12 2023
        ]
        
        for fmt in formatos:
            try:
                return datetime.strptime(date_str, fmt).strftime("%d/%m/%Y")
            except:
                continue
                
        # Se nenhum formato funcionar, tentar extrair os números
        import re
        numeros = re.findall(r'\d+', date_str)
        if len(numeros) >= 3:
            dia = numeros[0].zfill(2)
            mes = numeros[1].zfill(2)
            ano = numeros[2]
            if len(ano) == 2:
                ano = '20' + ano
            return f"{dia}/{mes}/{ano}"
            
        return ""
    except:
        return ""

def format_time(time_str):
    """Converte a hora para o formato HH:MM"""
    if not time_str or time_str == "Não informado":
        return ""
    try:
        # Remover possíveis textos extras
        time_str = time_str.lower().replace('h', ':').replace('hrs', ':00')
        
        # Tentar diferentes formatos de entrada
        formatos = [
            "%H:%M",     # 14:30
            "%H:%M:%S",  # 14:30:00
            "%I:%M %p",  # 02:30 PM
            "%H",        # 14
        ]
        
        for fmt in formatos:
            try:
                return datetime.strptime(time_str, fmt).strftime("%H:%M")
            except:
                continue
                
        # Se nenhum formato funcionar, tentar extrair os números
        import re
        numeros = re.findall(r'\d+', time_str)
        if len(numeros) >= 1:
            hora = numeros[0].zfill(2)
            minuto = numeros[1].zfill(2) if len(numeros) > 1 else "00"
            return f"{hora}:{minuto}"
            
        return ""
    except:
        return ""

def render_atraso_voo():
    """Renderiza a página de Atraso/Cancelamento de Voo"""
    st.title("Atraso / Cancelamento de Voo")
    
    # Inicializar dados do caso no session_state
    if 'selected_case_data' not in st.session_state:
        st.session_state.selected_case_data = {
            'assunto_caso': 'Atraso/Cancelamento de Voo',
            'pasta_caso_id': st.secrets.get("ROOT_FOLDER_ID")  # Usar pasta raiz como fallback
        }
    
    # Renderizar cada seção
    render_client_section()
    render_company_section()
    render_facts_section()

def generate_and_save_petition(st_session_state):
    """Gera e salva a petição preenchida na pasta do caso"""
    try:
        # Verificações detalhadas
        if 'selected_client_data' not in st_session_state:
            raise Exception("Dados do cliente não encontrados. Por favor, selecione um cliente.")
        
        if 'selected_company_data' not in st_session_state:
            raise Exception("Dados da empresa não encontrados. Por favor, selecione uma empresa.")
        
        if 'selected_case_data' not in st_session_state:
            raise Exception("Dados do caso não encontrados. Por favor, recarregue a página.")

        # Verificar se temos o ID da pasta
        pasta_caso_id = st_session_state.selected_case_data.get('pasta_caso_id')
        if not pasta_caso_id:
            pasta_caso_id = st.secrets.get("ROOT_FOLDER_ID")  # Usar pasta raiz como fallback
            if not pasta_caso_id:
                raise Exception("ID da pasta do caso não encontrado")
        
        # Criar serviço do Google Drive
        credentials_dict = json.loads(st.secrets["GOOGLE_CREDENTIALS"])
        creds = service_account.Credentials.from_service_account_info(
            credentials_dict,
            scopes=['https://www.googleapis.com/auth/drive']  # Escopo completo do Drive
        )
        drive_service = build('drive', 'v3', credentials=creds)

        # Baixar o template (usando get_media para arquivo .docx)
        request = drive_service.files().get_media(
            fileId=st.secrets["TEMPLATE_ATRASO_VOO_ID"]
        )
        
        template_content = io.BytesIO()
        downloader = MediaIoBaseDownload(template_content, request)
        done = False
        while done is False:
            status, done = downloader.next_chunk()

        # Carregar o documento com python-docx
        doc = Document(template_content)

        # Preparar os dados para substituição
        replace_dict = {
            'nome_completo': st_session_state.selected_client_data.get('nome_completo', ''),
            'nacionalidade': st_session_state.selected_client_data.get('nacionalidade', ''),
            'estado_civil': st_session_state.selected_client_data.get('estado_civil', ''),
            'profissao': st_session_state.selected_client_data.get('profissao', ''),
            'rg': st_session_state.selected_client_data.get('rg', ''),
            'cpf': st_session_state.selected_client_data.get('cpf', ''),
            'endereco': st_session_state.selected_client_data.get('endereco', ''),
            'bairro': st_session_state.selected_client_data.get('bairro', ''),
            'cidade': st_session_state.selected_client_data.get('cidade', ''),
            'cep': st_session_state.selected_client_data.get('cep', ''),
            'nome_empresa_re': st_session_state.selected_company_data.get('nome', ''),
            'cnpj_empresa_re': st_session_state.selected_company_data.get('cnpj', ''),
            'endereco_empresa_re': st_session_state.selected_company_data.get('endereco', ''),
            'dos_fatos': st_session_state.get('generated_facts', ''),
            'tempo_atraso': st_session_state.flight_info.get('tempo_atraso', ''),
            'valor_danos_material': st_session_state.flight_info.get('valor_total_custos', ''),
            'valor_danos_material_extenso': st_session_state.get('valor_danos_material_extenso', ''),
            'explicacao_danos_material': st_session_state.flight_info.get('descricao_custos', ''),
            'vara_civil': st_session_state.get('vara_civil', ''),
            'tribunal_jurisprudencia_deveres_transportador': st_session_state.get('tribunal_jurisprudencia_deveres_transportador', ''),
            'jurisprudencia_deveres_transportador': st_session_state.get('jurisprudencia_deveres_transportador', ''),
            'tribunal_jurisprudencia_da_inteligencia': st_session_state.get('tribunal_jurisprudencia_da_inteligencia', ''),
            'jurisprudencia_da_inteligencia': st_session_state.get('jurisprudencia_da_inteligencia', ''),
            'tribunal_jurisprudencia_da_responsabilidadea': st_session_state.get('tribunal_jurisprudencia_da_responsabilidadea', ''),
            'jurisprudencia_da_responsabilidadea': st_session_state.get('jurisprudencia_da_responsabilidadea', ''),
            'tribunal_jurisprudencia_dos_prejuizos': st_session_state.get('tribunal_jurisprudencia_dos_prejuizos', ''),
            'jurisprudencia_dos_prejuizos': st_session_state.get('jurisprudencia_dos_prejuizos', ''),
            'motivos_danos_moral': st_session_state.get('motivos_danos_morais', ''),
            'valor_dano_moral': st_session_state.get('valor_dano_moral', ''),
            'valor_dano_moral_extenso': st_session_state.get('valor_dano_moral_extenso', ''),
            'valor_dano_moral_material': st_session_state.get('valor_dano_moral_material', ''),
            'valor_dano_moral_material_extenso': st_session_state.get('valor_dano_moral_material_extenso', ''),
            'data_extenso': st_session_state.get('data_extenso', '')
        }

        # Substituir os placeholders no documento
        for paragraph in doc.paragraphs:
            for key, value in replace_dict.items():
                if f'{{{{{key}}}}}' in paragraph.text:
                    paragraph.text = paragraph.text.replace(f'{{{{{key}}}}}', str(value))

        # Salvar o documento modificado
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)

        # Nome do arquivo com data e hora
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"Petição Inicial - {timestamp}.docx"

        # Criar arquivo no Google Drive
        file_metadata = {
            'name': file_name,
            'parents': [pasta_caso_id],
            'mimeType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        }
        
        media = MediaIoBaseUpload(
            output, 
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            resumable=True
        )
        
        file = drive_service.files().create(
            body=file_metadata,
            media_body=media,
            fields='id, webViewLink'
        ).execute()

        return file.get('webViewLink')

    except Exception as e:
        logger.error(f"Erro ao gerar petição: {str(e)}")
        raise Exception(f"Erro ao gerar petição: {str(e)}") 